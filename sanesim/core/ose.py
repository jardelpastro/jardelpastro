"""Ordem de Serviço para Execução (OSE) — folha de planilha.

Cada OSE (entidade `OseSheet` do projeto) agrupa trechos e gera uma folha
no formato do modelo SANEPAR: cabeçalho com número/locação/folha de
cadastro, identificação da rua, tabela de estaqueamento contínuo (20 m)
ao longo dos trechos, observações e bloco de assinaturas.

Os dados hidráulicos vêm da simulação; os demais campos (ruas, número,
observações, responsáveis, gabarito da régua) são editáveis na aba OSEs.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import OseSheet, Project
from .simulation import PipeResult, SimulationResult

_THIN = Side(style="thin")
_MEDIUM = Side(style="medium")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOX = Border(left=_MEDIUM, right=_MEDIUM, top=_MEDIUM, bottom=_MEDIUM)
_HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEAD_FONT = Font(bold=True, size=9)
_CELL_FONT = Font(size=9)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADERS = ["Estaca", "Dist. entre\nPiquetes (m)", "Dist.\nAcumulada (m)",
           "Cota do\nTerreno (m)", "Declividade\n(m/m)",
           "Cota Geratriz\nInf. Coletor (m)", "Altura do\nGabarito (m)",
           "Cota Bordo Sup.\nda Régua (m)", "Altura da\nRégua (m)",
           "Diâmetro\n(mm)", "Prof. da\nVala (m)", "Recobrimento\n(m)",
           "Observações"]
WIDTHS = [8, 11, 11, 11, 11, 13, 10, 13, 10, 9, 10, 12, 28]
NCOLS = len(HEADERS)


class OseError(Exception):
    pass


@dataclass
class OseRow:
    """Uma linha (estaca) da planilha de uma OSE."""
    stake_label: str
    dist_prev: float
    dist_accum: float
    ground: float
    slope: float
    invert: float
    gauge: float          # altura do gabarito (régua)
    board: float          # cota do bordo superior da régua
    ruler: float          # altura da régua acima do terreno
    diameter_mm: int
    depth: float
    cover: float
    obs: str


def ose_pipe_results(project: Project, result: SimulationResult,
                     ose: OseSheet) -> list[PipeResult]:
    """Resultados dos trechos da OSE, na ordem definida na OSE."""
    by_name = {r.pipe: r for r in result.pipes}
    out = []
    for pid in ose.pipe_ids:
        pipe = project.pipe_by_id(pid)
        if pipe is None:
            continue
        r = by_name.get(pipe.name)
        if r is not None:
            out.append(r)
    return out


def split_runs(pipes: list[PipeResult]) -> list[list[PipeResult]]:
    """Divide os trechos da OSE em ramais contínuos.

    Um ramal é uma sequência em que o PV de jusante de um trecho é o de
    montante do seguinte; quando a sequência quebra (a OSE contém um
    ramal afluente, por exemplo), o estaqueamento reinicia.
    """
    runs: list[list[PipeResult]] = []
    for r in pipes:
        if runs and runs[-1][-1].downstream == r.upstream:
            runs[-1].append(r)
        else:
            runs.append([r])
    return runs


def build_ose_rows(project: Project, result: SimulationResult,
                   ose: OseSheet, step: float = 20.0) -> list[OseRow]:
    """Estaqueamento da OSE, contínuo dentro de cada ramal.

    Estacas a cada `step` m mais os pontos de PV (limites de trecho);
    terreno e coletor interpolados linearmente dentro de cada trecho.
    Ramais não consecutivos reiniciam o estaqueamento em 0.
    """
    pipes = ose_pipe_results(project, result, ose)
    rows: list[OseRow] = []
    for i, run in enumerate(split_runs(pipes)):
        run_rows = _run_rows(run, ose, step)
        if run_rows and len(pipes) > len(run):
            label = (f"Ramal {i + 1}: {run[0].upstream} → "
                     f"{run[-1].downstream}")
            first = run_rows[0]
            first.obs = f"{label}; {first.obs}" if first.obs else label
        rows.extend(run_rows)
    return rows


def _run_rows(pipes: list[PipeResult], ose: OseSheet,
              step: float) -> list[OseRow]:
    if not pipes:
        return []
    # segmentos com posição acumulada
    segments: list[tuple[float, float, PipeResult]] = []
    x = 0.0
    for r in pipes:
        segments.append((x, x + r.length, r))
        x += r.length
    total = x

    positions: set[float] = {0.0, total}
    p = step
    while p < total - 1e-6:
        positions.add(p)
        p += step
    for start, _end, _r in segments[1:]:
        # limite de trecho já coincidente com uma estaca não duplica linha
        if not any(abs(start - q) < 1e-6 for q in positions):
            positions.add(start)

    gauge = ose.gauge_height
    rows: list[OseRow] = []
    prev_x = 0.0
    stake = 0
    for pos in sorted(positions):
        # segmento que contém a posição (limite pertence ao trecho seguinte,
        # pois é o tubo que será assentado a partir dali)
        seg = None
        for start, end, r in segments:
            if start - 1e-6 <= pos < end - 1e-6:
                seg = (start, end, r)
                break
        if seg is None:                     # última posição (fim da OSE)
            seg = segments[-1]
        start, end, r = seg
        frac = (pos - start) / r.length if r.length > 0 else 0.0
        ground = r.ground_up + (r.ground_down - r.ground_up) * frac
        invert = r.invert_up - r.slope * (pos - start)
        board = invert + gauge
        obs_parts = []
        for s_start, _s_end, s_r in segments:
            if abs(pos - s_start) < 1e-6:
                obs_parts.append(f"PV {s_r.upstream}")
                if rows:
                    drop = s_r.drop_up
                    if drop > 1e-3:
                        obs_parts.append(f"degrau {drop:.3f} m")
        if abs(pos - total) < 1e-6:
            obs_parts.append(f"PV {segments[-1][2].downstream}")

        is_multiple = abs(pos / step - round(pos / step)) < 1e-6
        label = str(stake) if is_multiple else f"+{pos % step:.2f}"
        rows.append(OseRow(
            stake_label=label,
            dist_prev=pos - prev_x,
            dist_accum=pos,
            ground=ground,
            slope=r.slope,
            invert=invert,
            gauge=gauge,
            board=board,
            ruler=board - ground,
            diameter_mm=r.diameter_mm,
            depth=ground - invert,
            cover=ground - invert - r.diameter_mm / 1000.0,
            obs="; ".join(obs_parts),
        ))
        prev_x = pos
        if is_multiple:
            stake += 1
    return rows


def _cell(ws, row: int, col: int, value, fmt: str | None = None,
          bold: bool = False, align=_CENTER, border=_BORDER):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _HEAD_FONT if bold else _CELL_FONT
    c.border = border
    c.alignment = align
    if fmt:
        c.number_format = fmt
    return c


def _label_value(ws, row: int, col_label: int, col_value_end: int,
                 label: str, value: str):
    c = ws.cell(row=row, column=col_label,
                value=f"{label} {value}".strip())
    c.font = _CELL_FONT
    c.alignment = _LEFT
    ws.merge_cells(start_row=row, start_column=col_label,
                   end_row=row, end_column=col_value_end)


def export_oses(project: Project, result: SimulationResult, path: str,
                step: float = 20.0) -> None:
    """Grava um arquivo .xlsx com uma folha por OSE do projeto."""
    if not project.oses:
        raise OseError("Nenhuma OSE cadastrada. Use a aba OSEs (ou "
                       "'Gerar OSEs') antes de exportar.")
    wb = Workbook()
    wb.remove(wb.active)
    for ose in project.oses:
        _ose_sheet(wb, project, result, ose, step)
    wb.save(path)


def _ose_sheet(wb, project: Project, result: SimulationResult,
               ose: OseSheet, step: float):
    title = f"OSE {ose.number}" if ose.number else f"OSE s-n {ose.id[:4]}"
    ws = wb.create_sheet(title[:31])
    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    pipes = ose_pipe_results(project, result, ose)
    length = sum(r.length for r in pipes)
    diameters = " / ".join(sorted({str(r.diameter_mm) for r in pipes},
                                  key=lambda s: int(s)))
    materials = " / ".join(dict.fromkeys(r.material for r in pipes))

    # ------------------------------------------------- cabeçalho (modelo)
    third = NCOLS // 3
    _cell(ws, 1, 1, f"Número da O.S.E.: {ose.number}", bold=True,
          align=_LEFT, border=_BOX)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=third)
    _cell(ws, 1, third + 1, f"Locação: {ose.location}", bold=True,
          align=_LEFT, border=_BOX)
    ws.merge_cells(start_row=1, start_column=third + 1,
                   end_row=1, end_column=2 * third)
    _cell(ws, 1, 2 * third + 1,
          f"Nº da Folha de Cadastro: {ose.cadastre_sheet}", bold=True,
          align=_LEFT, border=_BOX)
    ws.merge_cells(start_row=1, start_column=2 * third + 1,
                   end_row=1, end_column=NCOLS)

    c = ws.cell(row=2, column=1, value="ORDEM DE SERVIÇO PARA EXECUÇÃO")
    c.font = Font(bold=True, size=14)
    c.alignment = _CENTER
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOLS)

    system = project.info.system or project.title
    c = ws.cell(row=3, column=1, value=system.upper())
    c.font = Font(bold=True, size=11)
    c.alignment = _CENTER
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NCOLS)

    city = ose.city or project.info.city
    _label_value(ws, 4, 1, 3, "Cidade:", city)
    _label_value(ws, 4, 4, 7, "Rua:", ose.street)
    _label_value(ws, 4, 8, 9, "Lado:", ose.side)
    _label_value(ws, 4, 10, 11, "Extensão:",
                 f"{length:,.2f} m".replace(",", "\x00")
                 .replace(".", ",").replace("\x00", "."))
    _label_value(ws, 4, 12, NCOLS, "Diâmetro:", diameters)
    _label_value(ws, 5, 1, 7,
                 "Coletor entre Rua:",
                 f"{ose.between_street}  e Rua: {ose.and_street}")
    _label_value(ws, 5, 8, NCOLS, "Material:", materials)
    _label_value(ws, 6, 1, 3, "Data:",
                 _dt.date.today().strftime("%d/%m/%Y"))
    _label_value(ws, 6, 4, NCOLS, "Projeto:", project.title)

    # ------------------------------------------------------- estaqueamento
    hrow = 8
    for col, label in enumerate(HEADERS, start=1):
        c = ws.cell(row=hrow, column=col, value=label)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.border = _BORDER
        c.alignment = _CENTER

    rows = build_ose_rows(project, result, ose, step)
    r_ = hrow + 1
    for row in rows:
        _cell(ws, r_, 1, row.stake_label)
        _cell(ws, r_, 2, round(row.dist_prev, 2), "#,##0.00")
        _cell(ws, r_, 3, round(row.dist_accum, 2), "#,##0.00")
        _cell(ws, r_, 4, round(row.ground, 3), "#,##0.000")
        _cell(ws, r_, 5, round(row.slope, 5), "0.00000")
        _cell(ws, r_, 6, round(row.invert, 3), "#,##0.000")
        _cell(ws, r_, 7, round(row.gauge, 3), "#,##0.000")
        _cell(ws, r_, 8, round(row.board, 3), "#,##0.000")
        _cell(ws, r_, 9, round(row.ruler, 3), "#,##0.000")
        _cell(ws, r_, 10, row.diameter_mm, "0")
        _cell(ws, r_, 11, round(row.depth, 3), "#,##0.000")
        _cell(ws, r_, 12, round(row.cover, 3), "#,##0.000")
        _cell(ws, r_, 13, row.obs, align=_LEFT)
        r_ += 1

    # ------------------------------------------------------- observações
    r_ += 1
    c = ws.cell(row=r_, column=1,
                value="(*) OBSERVAÇÃO: " + (ose.observations or ""))
    c.font = _CELL_FONT
    c.alignment = _LEFT
    ws.merge_cells(start_row=r_, start_column=1, end_row=r_,
                   end_column=NCOLS)

    # ------------------------------------------------------- assinaturas
    r_ += 2
    blocks = [("Proposição", ose.resp_proposal),
              ("Aprovação", ose.resp_approval),
              ("Liberação para Execução", ose.resp_release),
              ("Execução/Cadastramento", ose.resp_execution)]
    per = max(2, NCOLS // len(blocks))
    col = 1
    for label, name in blocks:
        end_col = min(col + per - 1, NCOLS)
        c = ws.cell(row=r_, column=col, value=label)
        c.font = _HEAD_FONT
        c.border = _BOX
        c.alignment = _CENTER
        ws.merge_cells(start_row=r_, start_column=col,
                       end_row=r_, end_column=end_col)
        c = ws.cell(row=r_ + 1, column=col, value=name or "")
        c.font = _CELL_FONT
        c.border = _BOX
        c.alignment = _CENTER
        ws.merge_cells(start_row=r_ + 1, start_column=col,
                       end_row=r_ + 1, end_column=end_col)
        col = end_col + 1

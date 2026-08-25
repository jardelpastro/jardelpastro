"""Geração do memorial de cálculo em Excel (padrão do modelo do usuário).

Abas: Capa, Trechos, Nós, Dimensionamento e Resultados. A aba Resultados
usa duas linhas por trecho (montante/jusante), como no memorial modelo e
nas planilhas de OSE.
"""

from __future__ import annotations

import datetime as _dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Project
from .simulation import SimulationResult

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
_VIOL_FILL = PatternFill("solid", fgColor="FFC7CE")
_TITLE_FONT = Font(bold=True, size=14)
_HEAD_FONT = Font(bold=True, size=9)
_CELL_FONT = Font(size=9)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _header(ws, row: int, labels: list[str], widths: list[float] | None = None):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = _HEAD_FONT
        c.fill = _HEAD_FILL
        c.border = _BORDER
        c.alignment = _CENTER
        if widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]


def _cell(ws, row: int, col: int, value, fmt: str | None = None,
          fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _CELL_FONT
    c.border = _BORDER
    c.alignment = _CENTER
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def _sheet_title(ws, title: str, project_title: str, ncols: int):
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT
    ws.cell(row=2, column=1, value=project_title).font = Font(bold=True)
    ws.cell(row=2, column=max(1, ncols - 1), value="DATA:").font = _HEAD_FONT
    d = ws.cell(row=2, column=ncols, value=_dt.date.today())
    d.number_format = "DD/MM/YYYY"


def export_memorial(project: Project, result: SimulationResult,
                    path: str) -> None:
    wb = Workbook()

    # ------------------------------------------------------------- Capa
    ws = wb.active
    ws.title = "Capa"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["F"].width = 32
    ws.column_dimensions["I"].width = 12
    ws["A1"] = project.title
    ws["A1"].font = _TITLE_FONT
    ws["A3"] = "Critérios de Dimensionamento"
    ws["A3"].font = Font(bold=True, size=12)

    c = project.criteria
    ws["A5"] = "Início de Plano"
    ws["A5"].font = Font(bold=True)
    ws["F5"] = "Fim de Plano"
    ws["F5"].font = Font(bold=True)
    start_rows = [
        ("População (hab) :", c.start.population),
        ("Consumo per capita (l/hab/dia) :", c.start.per_capita),
        ("Coeficiente de retorno :", c.start.return_coef),
        ("K2 :", c.start.k2),
        ("K3 :", c.start.k3),
    ]
    end_rows = [
        ("População (hab) :", c.end.population),
        ("Consumo per capita (l/hab/dia) :", c.end.per_capita),
        ("Coeficiente de retorno :", c.end.return_coef),
        ("K1 :", c.end.k1),
        ("K2 :", c.end.k2),
        ("K3 :", c.end.k3),
    ]
    for i, (label, value) in enumerate(start_rows, start=6):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=4, value=value)
    for i, (label, value) in enumerate(end_rows, start=6):
        ws.cell(row=i, column=6, value=label)
        ws.cell(row=i, column=9, value=value)

    d = project.design
    o = project.options
    ws["A13"] = "Gerais"
    ws["A13"].font = Font(bold=True)
    ws["F13"] = "Critério de Cálculo"
    ws["F13"].font = Font(bold=True)
    general_rows = [
        ("Vazão mínima (l/s) :", d.min_flow_lps),
        ("Diâmetro mínimo (mm) :", d.min_diameter_mm),
        ("Taxa de infiltração (l/s/km) :", c.infiltration_rate),
        ("Recobrimento mínimo (m) :", d.min_cover_m),
        ("Profundidade máxima (m) :", d.max_depth_m),
        ("Tensão trativa mínima (Pa) :", d.min_tractive_pa),
        ("Velocidade máxima (m/s) :", d.max_velocity_ms),
        ("Lâmina máxima (y/D) :", d.max_yd),
        ("Taxa linear inicial (l/s/km) :", round(result.rate_start, 3)),
        ("Taxa linear final (l/s/km) :", round(result.rate_end, 3)),
        ("Extensão total da rede (m) :", round(result.total_length_m, 2)),
    ]
    for i, (label, value) in enumerate(general_rows, start=14):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=4, value=value)
    n_txt = ("n do material" if o.use_material_n
             else f"n={o.n_fixed:.3f}".replace(".", ","))
    ws["F14"] = f"Manning  ({n_txt})"
    ws["F15"] = f"Modo de cálculo: {o.mode}"

    # ---------------------------------------------------------- Trechos
    ws = wb.create_sheet("Trechos")
    _sheet_title(ws, "TRECHOS", project.title, 8)
    _header(ws, 3, ["Rede", "Nó\nInicial", "Nó\nFinal", "Extensão\n(m)",
                    "Nome do\nTrecho", "Diâmetro\n(mm)", "Cota Mon\n(m)",
                    "Cota Jus\n(m)"],
            [8, 10, 10, 10, 10, 10, 10, 10])
    for i, r in enumerate(result.pipes, start=4):
        _cell(ws, i, 1, r.network)
        _cell(ws, i, 2, r.upstream)
        _cell(ws, i, 3, r.downstream)
        _cell(ws, i, 4, round(r.length, 2), "0.00")
        _cell(ws, i, 5, r.pipe)
        _cell(ws, i, 6, r.diameter_mm)
        _cell(ws, i, 7, round(r.invert_up, 3), "0.000")
        _cell(ws, i, 8, round(r.invert_down, 3), "0.000")

    # -------------------------------------------------------------- Nós
    ws = wb.create_sheet("Nós")
    _sheet_title(ws, "NÓS", project.title, 8)
    _header(ws, 3, ["Rede", "Nó", "Coordenada N\n(m)", "Coordenada E\n(m)",
                    "Cota Terreno\n(m)", "Q Pontual Ini\n(l/s)",
                    "Q Pontual Fim\n(l/s)", "Tipo\ndo Nó"],
            [8, 10, 14, 14, 12, 12, 12, 10])
    for i, node in enumerate(project.nodes, start=4):
        name = result.renamed.get(node.name, node.name)
        _cell(ws, i, 1, node.network)
        _cell(ws, i, 2, name)
        _cell(ws, i, 3, node.coord_n, "0.00")
        _cell(ws, i, 4, node.coord_e, "0.00")
        _cell(ws, i, 5, node.ground_elev, "0.000")
        _cell(ws, i, 6, node.q_point_start, "0.00")
        _cell(ws, i, 7, node.q_point_end, "0.00")
        _cell(ws, i, 8, node.node_type)

    # -------------------------------------------------- Dimensionamento
    ws = wb.create_sheet("Dimensionamento")
    _sheet_title(ws, "DIMENSIONAMENTO", project.title, 9)
    _header(ws, 3, ["Rede", "Nó\nInicial", "Nó\nFinal", "Extensão\n(m)",
                    "Nome do\nTrecho", "Recobr. Mín\n(m)", "Prof. Máx\n(m)",
                    "Situação", "Critério"],
            [8, 10, 10, 10, 10, 11, 10, 16, 10])
    pipes_by_name = {p.name: p for p in project.pipes}
    for i, r in enumerate(result.pipes, start=4):
        p = pipes_by_name.get(r.pipe)
        _cell(ws, i, 1, r.network)
        _cell(ws, i, 2, r.upstream)
        _cell(ws, i, 3, r.downstream)
        _cell(ws, i, 4, round(r.length, 2), "0.00")
        _cell(ws, i, 5, r.pipe)
        _cell(ws, i, 6, project.design.min_cover_m, "0.00")
        _cell(ws, i, 7, project.design.max_depth_m, "0.00")
        _cell(ws, i, 8, p.status if p else "Rede Projetada")
        _cell(ws, i, 9, "Global")

    # ------------------------------------------------------- Resultados
    ws = wb.create_sheet("Resultados")
    _sheet_title(ws, "RESULTADOS", project.title, 20)
    _header(ws, 3, ["Rede", "Trecho", "PV Inicial\nPV Fim", "Extensão\n(m)",
                    "Cont. Lin\n(l/s/km)\nini/fim",
                    "Cont. Trecho\n(l/s)\nini/fim", "Q Pontual\n(l/s)",
                    "Q Mont. (l/s)\nini/fim", "Q Jus (l/s)\nini/fim",
                    "Diâm.\n(mm)", "Decliv.\n(m/m)", "Cota Ter.\n(m)",
                    "Cota GI Col\n(m)", "Rec. Col. (m)\nmon/jus",
                    "Prof. Vala (m)\nmon/jus", "y/D\nini/fim",
                    "V (m/s)\nini/fim", "Arr. In. (Pa)\nVc (m/s)",
                    "n\nManning", "Larg. Vala\n(m)"],
            [7, 8, 9, 9, 9, 10, 9, 10, 10, 8, 9, 9, 10, 10, 10, 8, 8, 10,
             8, 9])
    row = 4
    for r in result.pipes:
        fill = _VIOL_FILL if r.violations else None
        # linha de montante / valores iniciais
        _cell(ws, row, 1, r.network, fill=fill)
        _cell(ws, row, 2, r.pipe, fill=fill)
        _cell(ws, row, 3, r.upstream, fill=fill)
        _cell(ws, row, 4, round(r.length, 2), "0.00", fill)
        _cell(ws, row, 5, round(r.rate_start, 2), "0.00", fill)
        _cell(ws, row, 6, round(r.q_reach_start, 3), "0.000", fill)
        _cell(ws, row, 7, round(r.q_point, 2), "0.00", fill)
        _cell(ws, row, 8, round(r.q_up_start, 3), "0.000", fill)
        _cell(ws, row, 9, round(r.q_down_start, 3), "0.000", fill)
        _cell(ws, row, 10, r.diameter_mm, fill=fill)
        _cell(ws, row, 11, round(r.slope, 4), "0.0000", fill)
        _cell(ws, row, 12, round(r.ground_up, 3), "0.000", fill)
        _cell(ws, row, 13, round(r.invert_up, 3), "0.000", fill)
        _cell(ws, row, 14, round(r.cover_up, 3), "0.000", fill)
        _cell(ws, row, 15, round(r.depth_up, 3), "0.000", fill)
        _cell(ws, row, 16, round(r.yd_start, 2), "0.00", fill)
        _cell(ws, row, 17, round(r.v_start, 2), "0.00", fill)
        _cell(ws, row, 18, round(r.tractive_pa, 2), "0.00", fill)
        _cell(ws, row, 19, r.n_manning, "0.000", fill)
        _cell(ws, row, 20, round(r.trench_width, 2), "0.00", fill)
        # linha de jusante / valores finais
        row += 1
        _cell(ws, row, 1, "", fill=fill)
        _cell(ws, row, 2, "", fill=fill)
        _cell(ws, row, 3, r.downstream, fill=fill)
        _cell(ws, row, 4, "", fill=fill)
        _cell(ws, row, 5, round(r.rate_end, 2), "0.00", fill)
        _cell(ws, row, 6, round(r.q_reach_end, 3), "0.000", fill)
        _cell(ws, row, 7, "", fill=fill)
        _cell(ws, row, 8, round(r.q_up_end, 3), "0.000", fill)
        _cell(ws, row, 9, round(r.q_down_end, 3), "0.000", fill)
        _cell(ws, row, 10, "", fill=fill)
        _cell(ws, row, 11, "", fill=fill)
        _cell(ws, row, 12, round(r.ground_down, 3), "0.000", fill)
        _cell(ws, row, 13, round(r.invert_down, 3), "0.000", fill)
        _cell(ws, row, 14, round(r.cover_down, 3), "0.000", fill)
        _cell(ws, row, 15, round(r.depth_down, 3), "0.000", fill)
        _cell(ws, row, 16, round(r.yd_end, 2), "0.00", fill)
        _cell(ws, row, 17, round(r.v_end, 2), "0.00", fill)
        _cell(ws, row, 18, round(r.v_critical, 2), "0.00", fill)
        _cell(ws, row, 19, r.n_manning, "0.000", fill)
        _cell(ws, row, 20, "", fill=fill)
        row += 1

    # Observações de violação ao final
    if result.has_violations:
        row += 1
        ws.cell(row=row, column=1, value="OBSERVAÇÕES / VIOLAÇÕES:").font = \
            Font(bold=True)
        for r in result.pipes:
            for v in r.violations:
                row += 1
                ws.cell(row=row, column=1, value=f"{r.pipe}: {v}").font = \
                    _CELL_FONT

    wb.save(path)

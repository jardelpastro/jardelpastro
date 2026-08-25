import openpyxl
import pytest

from sanesim.core.models import OseSheet, Project
from sanesim.core.ose import OseError, build_ose_rows, export_oses
from sanesim.core.simulation import simulate


@pytest.fixture()
def project_result():
    proj = Project.load("examples/exemplo_rede.json")
    res = simulate(proj)
    return proj, res


def test_ensure_oses_groups_by_network(project_result):
    proj, _res = project_result
    proj.oses = []
    created = proj.ensure_oses()
    assert created == 1                      # exemplo tem uma rede (C1)
    assert len(proj.oses[0].pipe_ids) == len(proj.pipes)
    # idempotente: trechos já cobertos não geram novas OSEs
    assert proj.ensure_oses() == 0


def test_build_ose_rows_continuous_stationing(project_result):
    proj, res = project_result
    # OSE com os dois primeiros trechos (T1 e T2, consecutivos)
    ose = OseSheet(number="10", pipe_ids=[proj.pipes[0].id,
                                          proj.pipes[1].id])
    rows = build_ose_rows(proj, res, ose)
    t1, t2 = res.pipes[0], res.pipes[1]
    total = t1.length + t2.length
    assert rows[0].dist_accum == 0.0
    assert rows[-1].dist_accum == pytest.approx(total)
    # o limite entre T1 e T2 aparece como estaca intermediária com PV
    boundary = [r for r in rows
                if abs(r.dist_accum - t1.length) < 1e-6]
    assert boundary and "PV" in boundary[0].obs
    # a distância acumulada é estritamente crescente
    accum = [r.dist_accum for r in rows]
    assert accum == sorted(accum)
    # régua: bordo = geratriz inferior + gabarito
    for r in rows:
        assert r.board == pytest.approx(r.invert + ose.gauge_height)
        assert r.depth - r.cover == pytest.approx(r.diameter_mm / 1000.0)


def test_gauge_height_adjustable(project_result):
    proj, res = project_result
    ose = OseSheet(number="1", pipe_ids=[proj.pipes[0].id],
                   gauge_height=4.5)
    rows = build_ose_rows(proj, res, ose)
    assert all(r.gauge == pytest.approx(4.5) for r in rows)
    assert rows[0].board == pytest.approx(rows[0].invert + 4.5)


def test_export_oses_one_sheet_per_ose(project_result, tmp_path):
    proj, res = project_result
    proj.oses = []
    proj.ensure_oses()
    ose = proj.oses[0]
    ose.number = "142"
    ose.street = "RUA JULIO MANFREDINI JUNIOR"
    ose.side = "ESQUERDO"
    ose.observations = "Executar com escoramento contínuo."
    ose.resp_proposal = "Eng. Fulano"
    out = tmp_path / "oses.xlsx"
    export_oses(proj, res, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["OSE 142"]
    ws = wb["OSE 142"]
    text = " ".join(str(c.value) for row in ws.iter_rows()
                    for c in row if c.value)
    assert "ORDEM DE SERVIÇO PARA EXECUÇÃO" in text
    assert "Número da O.S.E.: 142" in text
    assert "RUA JULIO MANFREDINI JUNIOR" in text
    assert "ESQUERDO" in text
    assert "escoramento contínuo" in text
    assert "Proposição" in text and "Eng. Fulano" in text
    assert "Liberação para Execução" in text


def test_export_without_oses_raises(project_result, tmp_path):
    proj, res = project_result
    proj.oses = []
    with pytest.raises(OseError):
        export_oses(proj, res, str(tmp_path / "x.xlsx"))


def test_ose_serialization_roundtrip(project_result, tmp_path):
    proj, _res = project_result
    proj.oses = []
    proj.ensure_oses()
    proj.oses[0].number = "77"
    proj.oses[0].gauge_height = 5.0
    proj.info.city = "Campo Largo"
    proj.info.designer = "Eng. Jardel"
    path = tmp_path / "p.json"
    proj.save(str(path))
    loaded = Project.load(str(path))
    assert loaded.oses[0].number == "77"
    assert loaded.oses[0].gauge_height == 5.0
    assert loaded.oses[0].pipe_ids == proj.oses[0].pipe_ids
    assert loaded.info.city == "Campo Largo"
    assert loaded.info.designer == "Eng. Jardel"

"""Tela principal: traçado contínuo, invalidação e gestão de OSEs."""

import os

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

pytest.importorskip("PySide6")

from PySide6.QtCore import QPointF  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from sanesim.core.models import OseSheet, Project  # noqa: E402
from sanesim.ui.main_window import MainWindow  # noqa: E402


@pytest.fixture(scope="module")
def app():
    yield QApplication.instance() or QApplication([])


@pytest.fixture()
def window(app):
    w = MainWindow()
    w.project = Project.load("examples/exemplo_rede.json")
    w._load_all()
    return w


def test_draw_mode_chains_nodes_and_pipes(window):
    plan = window.plan_tab
    plan.set_mode("draw")
    n_nodes = len(window.project.nodes)
    n_pipes = len(window.project.pipes)
    # três cliques no vazio: 3 PVs e 2 trechos encadeados
    plan.draw_click(QPointF(672300.0, -7184700.0))
    plan.draw_click(QPointF(672350.0, -7184650.0))
    plan.draw_click(QPointF(672400.0, -7184600.0))
    assert len(window.project.nodes) == n_nodes + 3
    assert len(window.project.pipes) == n_pipes + 2
    p1, p2 = window.project.pipes[-2], window.project.pipes[-1]
    assert p1.downstream == p2.upstream
    assert p2.upstream == window.project.nodes[-2].name


def test_draw_mode_connects_to_existing_node(window):
    plan = window.plan_tab
    plan.set_mode("draw")
    existing = window.project.node_by_name("PV-01")
    n_nodes = len(window.project.nodes)
    # clica em cima de um PV existente: inicia dele, sem criar nó
    plan.draw_click(QPointF(existing.coord_e, -existing.coord_n))
    assert len(window.project.nodes) == n_nodes
    plan.draw_click(QPointF(existing.coord_e + 80.0,
                            -existing.coord_n))
    assert window.project.pipes[-1].upstream == "PV-01"


def test_draw_eee_ends_drawing(window):
    plan = window.plan_tab
    plan.set_mode("draw")
    plan.draw_click(QPointF(672600.0, -7184500.0))
    plan.node_type_combo.setCurrentText("EEE")
    plan.draw_click(QPointF(672650.0, -7184450.0))
    new_node = window.project.nodes[-1]
    assert new_node.node_type == "EEE"
    assert plan.mode == "select"          # elevatória encerra o traçado
    assert window.project.pipes[-1].downstream == new_node.name


def test_network_change_invalidates_results(window):
    window.run_simulation()
    assert window.last_result is not None
    window.plan_tab.add_node_at(QPointF(672700.0, -7184400.0))
    assert window.last_result is None     # topologia mudou: descarta


def test_live_update_during_drag(window):
    plan = window.plan_tab
    node = window.project.nodes[0]
    item = plan._node_items[node.id]
    item.setSelected(True)
    item.setPos(672123.0, -7184987.0)     # simula o arrasto
    assert "672.123" in plan.n_coord_e.text()
    assert "7.184.987" in plan.n_coord_n.text()


def test_ose_sequential_numbering(window):
    ose_tab = window.ose_tab
    window.project.oses = [OseSheet(number="41")]
    ose_tab.load_from(window.project)
    ose_tab._new_ose()
    assert window.project.oses[-1].number == "42"
    ose_tab._new_ose()
    assert window.project.oses[-1].number == "43"


def test_ose_cancelada_flag_and_serialization(tmp_path):
    proj = Project.load("examples/exemplo_rede.json")
    proj.ensure_oses()
    proj.oses[0].status = "cancelada"
    path = tmp_path / "c.json"
    proj.save(str(path))
    loaded = Project.load(str(path))
    assert loaded.oses[0].status == "cancelada"


def test_cancelada_marked_in_export(tmp_path):
    import openpyxl
    from sanesim.core.ose import export_oses
    from sanesim.core.simulation import simulate
    proj = Project.load("examples/exemplo_rede.json")
    proj.oses = []
    proj.ensure_oses()
    proj.oses[0].number = "7"
    proj.oses[0].status = "cancelada"
    res = simulate(proj)
    out = tmp_path / "oses.xlsx"
    export_oses(proj, res, str(out))
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == ["OSE 7 CANCELADA"]
    text = " ".join(str(c.value) for row in wb.active.iter_rows()
                    for c in row if c.value)
    assert "CANCELADA" in text
